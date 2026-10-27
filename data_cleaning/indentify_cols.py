import pandas as pd
from openpyxl import load_workbook

merged = pd.read_excel("USI_FULL_Database.xlsx")
sheet_cols = set(pd.read_excel("data_cleaning/data_for_cleaning/PROJECT SHEETS MASTER LIST - with coordinates.xlsx", sheet_name="master list").columns.str.strip())

wb = load_workbook("data_cleaning/data_for_cleaning/USI_Toronto_Projects_Emmett.xlsx")
table_cols = set()
for ws in wb.worksheets:
    for t in ws.tables.values():
        table_cols |= {c.value.strip() for c in ws[t.ref][0]}

for c in merged.columns:
    src = "BOTH" if c.endswith(("_table", "_sheet")) or c == "PROJECT #" else \
          "TABLE_FILE" if c in table_cols else "SHEET_FILE" if c in sheet_cols else "?"
    print(f"{src:12} {c}")