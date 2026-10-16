import pandas as pd
from openpyxl import load_workbook

# ---- Update these ----
TABLE_FILE = "data_cleaning/data_for_cleaning/USI_Toronto_Projects_Emmett.xlsx"   # workbook containing the formatted Excel table
TABLE_SHEET = "Join Sheet"                        # sheet the table is on, or None to search every sheet
TABLE_NAME = None                     # table name (Table Design tab), or None to use the first table found
SHEET_FILE = "data_cleaning/data_for_cleaning/PROJECT SHEETS MASTER LIST - with coordinates.xlsx"   # workbook with the plain (unformatted) data
SHEET_NAME = "master list"                 # exact tab name of the plain data
OUTPUT_FILE = "merged.xlsx"
KEY = "PROJECT #"
HOW = "right"                         # "inner" = matches only, "left"/"outer" = keep unmatched rows
# ----------------------


def read_excel_table(path, sheet, table_name=None):
    """Read a formatted Excel table wherever it sits in the workbook."""
    wb = load_workbook(path, data_only=True)
    sheets = [wb[sheet]] if sheet else wb.worksheets
    for ws in sheets:
        for name, ref in ws.tables.items():  # items() gives (name, "B3:C6")
            if table_name in (None, name):
                print(f"Using table '{name}' on sheet '{ws.title}'")
                rows = [[c.value for c in row] for row in ws[ref]]
                return pd.DataFrame(rows[1:], columns=rows[0])
    raise ValueError("Formatted table not found - check TABLE_SHEET / TABLE_NAME")


def clean(df):
    df.columns = [str(c).strip() for c in df.columns]           # fix stray spaces in headers
    df[KEY] = df[KEY].astype(str).str.strip().str.removesuffix(".0")  # 101 vs "101" vs 101.0
    return df


table_df = clean(read_excel_table(TABLE_FILE, TABLE_SHEET, TABLE_NAME))
sheet_df = clean(pd.read_excel(SHEET_FILE, sheet_name=SHEET_NAME))

merged = table_df.merge(sheet_df, on=KEY, how=HOW, suffixes=("_table", "_sheet"))
merged.to_excel(OUTPUT_FILE, sheet_name="Merged", index=False)

print(f"Table rows: {len(table_df)} | Sheet rows: {len(sheet_df)} | Merged rows: {len(merged)}")